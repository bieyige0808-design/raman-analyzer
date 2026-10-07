"""Multi-file Raman fitting and editable Excel export.

Run from the project root:
    .venv/bin/python raman_batch.py
Defaults preserve the validated notebook workflow.
"""
from pathlib import Path
from datetime import datetime
import argparse

import numpy as np
import pandas as pd
from scipy import sparse
from scipy.sparse.linalg import spsolve
from lmfit.models import LorentzianModel, ConstantModel

PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_DIR = PROJECT_DIR / "data" / "raw"
DEFAULT_EXCEL_DIR = Path.home() / "Desktop"
FIT_MIN = 1000
FIT_MAX = 3000
BASELINE_LAM = 1e5
BASELINE_P = 0.01

def baseline_als(y, lam=1e5, p=0.01, niter=10):
    y = np.asarray(y, dtype=float)
    length = len(y)
    difference = sparse.diags([1, -2, 1], [0, 1, 2], shape=(length - 2, length), dtype=float)
    weights = np.ones(length)

    for _ in range(niter):
        weight_matrix = sparse.spdiags(weights, 0, length, length)
        system = weight_matrix + lam * difference.T @ difference
        baseline = spsolve(system.tocsc(), weights * y)
        weights = np.where(y > baseline, p, 1 - p)

    return baseline


def build_model():
    background = ConstantModel(prefix="bkg_")
    d_peak = LorentzianModel(prefix="d_")
    g_peak = LorentzianModel(prefix="g_")
    two_d_peak = LorentzianModel(prefix="two_d_")

    model = background + d_peak + g_peak + two_d_peak

    initial_params = model.make_params()

    initial_params["bkg_c"].set(value=0)

    initial_params["d_center"].set(
        value=1350, min=1250, max=1450
    )
    initial_params["d_sigma"].set(
        value=40, min=1, max=200
    )
    initial_params["d_amplitude"].set(value=1000, min=0)

    initial_params["g_center"].set(
        value=1580, min=1500, max=1650
    )
    initial_params["g_sigma"].set(
        value=30, min=1, max=200
    )
    initial_params["g_amplitude"].set(value=1000, min=0)

    initial_params["two_d_center"].set(
        value=2700, min=2550, max=2800
    )
    initial_params["two_d_sigma"].set(
        value=40, min=1, max=200
    )
    initial_params["two_d_amplitude"].set(value=1000, min=0)
    return model, initial_params

def analyze_one_file(file_path, model, initial_params, all_curve_data,
                     fit_min, fit_max, baseline_lam, baseline_p, subtract_baseline=True):
    df = pd.read_csv(
        file_path,
        sep="\t",
        skiprows=1,
        header=None,
    )

    if df.shape[1] != 2:
        raise ValueError(
            f"Expected two columns; found {df.shape[1]}."
        )

    df.columns = ["raman_shift", "intensity"]

    df = df.apply(pd.to_numeric, errors="coerce")
    df = df.replace([np.inf, -np.inf], np.nan).dropna()
    df = df.sort_values("raman_shift")

    region = df[
        df["raman_shift"].between(fit_min, fit_max)
    ]

    if len(region) < 20:
        raise ValueError("Fewer than 20 points in the fitting range.")

    x = region["raman_shift"].to_numpy()
    y = region["intensity"].to_numpy()

    baseline = baseline_als(
        y,
        lam=baseline_lam,
        p=baseline_p,
    ) if subtract_baseline else np.zeros_like(y, dtype=float)
    corrected = y - baseline

    if not np.isfinite(corrected).all():
        raise ValueError("Baseline correction produced non-finite values.")

    if corrected.max() <= 0:
        raise ValueError("No positive signal after baseline correction.")

    fit_params = initial_params.copy()

    fit_params["d_amplitude"].set(
        value=float(corrected.max()) * 100
    )
    fit_params["g_amplitude"].set(
        value=float(corrected.max()) * 100
    )
    two_d_mask = (x >= 2550) & (x <= 2800)

    if two_d_mask.sum() < 5:
        raise ValueError("Insufficient data in the 2D peak region.")

    two_d_height = max(
        float(corrected[two_d_mask].max()),
        1e-6,
    )

    fit_params["two_d_amplitude"].set(
        value=two_d_height * np.pi * 40
    )

    result = model.fit(
        corrected,
        fit_params,
        x=x,
    )

    if not result.success:
        raise RuntimeError(f"Fit did not converge: {result.message}")

    components = result.eval_components(x=x)
    residual = corrected - result.best_fit

    d_area = result.params["d_amplitude"].value
    g_area = result.params["g_amplitude"].value

    if g_area <= 0:
        raise ValueError("G peak area must be positive.")

    curves = pd.DataFrame({
        "raman_shift": x,
        "raw_intensity": y,
        "baseline": baseline,
        "corrected_intensity": corrected,
        "total_fit": result.best_fit,
        "d_peak": components["d_"],
        "g_peak": components["g_"],
        "2D_peak": components["two_d_"],
        "residual": residual,
    })

    all_curve_data[file_path.name] = {
        "raw": df.copy(),
        "fitted": curves.copy(),
    }


    return {
        "file": file_path.name,
        "status": "success",
        "baseline_subtracted": "Yes" if subtract_baseline else "No",
        "I_D": result.params["d_height"].value,
        "I_G": result.params["g_height"].value,
        "I_D/I_G": (
            result.params["d_height"].value
            / result.params["g_height"].value
        ),
        "D_center": result.params["d_center"].value,
        "G_center": result.params["g_center"].value,
        "D_FWHM": result.params["d_fwhm"].value,
        "G_FWHM": result.params["g_fwhm"].value,
        "D_area": d_area,
        "G_area": g_area,
        "D_G_area_ratio": d_area / g_area,
        "2D_center": result.params["two_d_center"].value,
        "2D_FWHM": result.params["two_d_fwhm"].value,
        "2D_area": result.params["two_d_amplitude"].value,
        "2D_G_area_ratio": (
            result.params["two_d_amplitude"].value / g_area
        ),
        "residual_RMSE": float(
            np.sqrt(np.mean(residual ** 2))
        ),
        "error": "",
    }




def export_workbook(records, all_curve_data, files, excel_path, fit_min, fit_max):
    from datetime import datetime
    from pathlib import Path

    import numpy as np
    import pandas as pd

    from openpyxl import Workbook
    from openpyxl.chart import ScatterChart, Reference, Series
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter


    # ---------- Workbook settings ----------

    wb = Workbook()

    overview_ws = wb.active
    overview_ws.title = "Raw_Overview"

    summary_ws = wb.create_sheet("Summary")

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="24476A",
    )
    average_fill = PatternFill(
        fill_type="solid",
        fgColor="E2EFDA",
    )
    header_font = Font(
        name="Calibri",
        size=11,
        bold=True,
        color="FFFFFF",
    )
    body_font = Font(name="Calibri", size=11)

    colors = [
        "1F77B4",
        "FF7F0E",
        "2CA02C",
        "D62728",
        "9467BD",
        "8C564B",
        "E377C2",
        "7F7F7F",
        "BCBD22",
        "17BECF",
    ]


    # ---------- Helper functions ----------

    def excel_value(value):
        if value is None:
            return None

        if isinstance(value, np.generic):
            value = value.item()

        if isinstance(value, float) and not np.isfinite(value):
            return None

        return value


    def write_dataframe(ws, frame, header_row, start_col):
        for offset, column in enumerate(frame.columns):
            cell = ws.cell(
                row=header_row,
                column=start_col + offset,
                value=str(column),
            )
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        ws.row_dimensions[header_row].height = 32

        for row_offset, values in enumerate(
            frame.itertuples(index=False, name=None),
            start=1,
        ):
            for col_offset, value in enumerate(values):
                cell = ws.cell(
                    row=header_row + row_offset,
                    column=start_col + col_offset,
                    value=excel_value(value),
                )
                cell.font = body_font

                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0000"


    def create_chart(title, y_title="Intensity (a.u.)"):
        chart = ScatterChart()
        chart.title = title
        chart.scatterStyle = "line"
        chart.x_axis.title = "Raman Shift (cm^-1)"
        chart.y_axis.title = y_title
        chart.width = 24
        chart.height = 11
        chart.legend.position = "b"
        chart.visible_cells_only = False

        return chart


    def add_series(
        chart,
        ws,
        x_col,
        y_col,
        first_row,
        last_row,
        title,
        color,
    ):
        x_values = Reference(
            ws,
            min_col=x_col,
            min_row=first_row,
            max_row=last_row,
        )
        y_values = Reference(
            ws,
            min_col=y_col,
            min_row=first_row,
            max_row=last_row,
        )

        series = Series(
            y_values,
            x_values,
            title=title,
        )
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 18000
        series.marker.symbol = "none"
        series.smooth = False

        chart.series.append(series)


    # ---------- English sample identifiers ----------

    file_ids = {file_path.name: file_path.stem for file_path in files}

    summary = pd.DataFrame(records)

    if summary.empty:
        raise RuntimeError("No results are available for export.")

    # Export curves only for samples reported as successful.
    successful_names = set(
        summary.loc[summary["status"] == "success", "file"]
    )

    sample_names = [
        file_path.name
        for file_path in files
        if (
            file_path.name in successful_names
            and file_path.name in all_curve_data
        )
    ]


    # ---------- Sheet 1: Summary ----------

    export_summary = summary.copy()

    export_summary["file"] = export_summary["file"].map(file_ids)
    export_summary = export_summary.rename(
        columns={"file": "Sample_ID"}
    )

    # Keep the three primary metrics immediately after Sample_ID.
    primary_columns = [
        "status",
        "Sample_ID",
        "I_D",
        "I_G",
        "I_D/I_G",
    ]

    for column in primary_columns:
        if column not in export_summary.columns:
            export_summary[column] = np.nan

    remaining_columns = [
        column
        for column in export_summary.columns
        if column not in primary_columns
    ]

    export_summary = export_summary[
        primary_columns + remaining_columns
    ]

    # Row 1 contains headers.
    # Row 2 contains averages.
    # Sample results begin in row 3.
    average_record = {
        column: None
        for column in export_summary.columns
    }
    average_record["Sample_ID"] = "Average"

    summary_with_average = pd.concat(
        [
            pd.DataFrame([average_record]),
            export_summary,
        ],
        ignore_index=True,
    )

    write_dataframe(
        summary_ws,
        summary_with_average,
        header_row=1,
        start_col=1,
    )

    status_col = (
        export_summary.columns.get_loc("status") + 1
    )
    status_letter = get_column_letter(status_col)

    last_summary_row = len(export_summary) + 2

    # Average each numeric metric over successful samples.
    # The ratio column averages individual sample ratios.
    for column_index, column_name in enumerate(
        export_summary.columns,
        start=1,
    ):
        if pd.api.types.is_numeric_dtype(
            export_summary[column_name]
        ):
            column_letter = get_column_letter(column_index)

            summary_ws.cell(
                row=2,
                column=column_index,
                value=(
                    f'=IFERROR(AVERAGEIF('
                    f'${status_letter}$3:'
                    f'${status_letter}${last_summary_row},'
                    f'"success",'
                    f'{column_letter}3:'
                    f'{column_letter}{last_summary_row}'
                    f'),"")'
                ),
            ).number_format = "0.0000"

    for cell in summary_ws[2]:
        cell.fill = average_fill
        cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
        )

    summary_ws.freeze_panes = "B3"

    for column_index, column_name in enumerate(
        export_summary.columns,
        start=1,
    ):
        letter = get_column_letter(column_index)
        summary_ws.column_dimensions[letter].width = (
            36 if column_name == "error" else 22
        )


    # ---------- Sheet 2: Normalized raw overview ----------

    overview_chart = create_chart(
        "Normalized Raw Spectra with Vertical Offsets",
        y_title="Normalized Intensity + Offset",
    )
    overview_chart.height = 15

    # The chart occupies the upper-left area.
    # Data blocks start below the chart and extend horizontally.
    overview_ws.add_chart(overview_chart, "A1")

    overview_title_row = 33
    overview_header_row = 34
    overview_first_row = 35

    for sample_index, source_name in enumerate(sample_names):
        sample_id = file_ids[source_name]
        raw = all_curve_data[source_name]["raw"].copy()

        max_intensity = float(raw["intensity"].max())

        raw["Normalized_Intensity"] = np.nan
        raw["Offset_Intensity"] = np.nan

        valid_normalization = (
            np.isfinite(max_intensity)
            and max_intensity > 0
        )

        if valid_normalization:
            raw["Normalized_Intensity"] = (
                raw["intensity"] / max_intensity
            )
            raw["Offset_Intensity"] = (
                raw["Normalized_Intensity"] + sample_index
            )
        else:
            print(
                f"{sample_id}: normalization skipped "
                "because the maximum intensity is not positive."
            )

        overview_data = pd.DataFrame({
            "Raman_Shift": raw["raman_shift"],
            "Raw_Intensity": raw["intensity"],
            "Normalized_Intensity": raw["Normalized_Intensity"],
            "Offset_Intensity": raw["Offset_Intensity"],
        })

        # Four data columns plus one blank separator column.
        start_col = 1 + sample_index * 5

        overview_ws.cell(
            row=overview_title_row,
            column=start_col,
            value=f"{sample_id} | Offset = {sample_index}",
        ).font = Font(
            name="Calibri",
            size=12,
            bold=True,
        )

        write_dataframe(
            overview_ws,
            overview_data,
            header_row=overview_header_row,
            start_col=start_col,
        )

        for column_index in range(start_col, start_col + 4):
            overview_ws.column_dimensions[
                get_column_letter(column_index)
            ].width = 23

        overview_ws.column_dimensions[
            get_column_letter(start_col + 4)
        ].width = 4

        if valid_normalization:
            add_series(
                overview_chart,
                overview_ws,
                x_col=start_col,
                y_col=start_col + 3,
                first_row=overview_first_row,
                last_row=overview_header_row + len(overview_data),
                title=sample_id,
                color=colors[sample_index % len(colors)],
            )


    # ---------- Individual sample sheets ----------

    for source_name in sample_names:
        sample_id = file_ids[source_name]
        sample_data = all_curve_data[source_name]

        raw = sample_data["raw"].copy()
        fitted = sample_data["fitted"].copy()

        sheet_name = sample_id
        for character in ['\\\\', '/', '*', '?', ':', '[', ']']:
            sheet_name = sheet_name.replace(character, "_")
        sheet_name = sheet_name[:31].strip("'") or "Data"
        if sheet_name.lower() == "history":
            sheet_name = "History_Data"
        base_name = sheet_name
        number = 2
        while sheet_name.lower() in {name.lower() for name in wb.sheetnames}:
            suffix = f"_{number}"
            sheet_name = base_name[:31 - len(suffix)] + suffix
            number += 1
        ws = wb.create_sheet(sheet_name)

        # Three editable charts at the upper-left.
        # Numerical data starts below all three charts.
        data_title_row = 83
        data_header_row = 84
        first_data_row = 85

        raw_export = raw.rename(columns={
            "raman_shift": "Raw_Raman_Shift",
            "intensity": "Raw_Intensity",
        })

        ws.cell(
            row=data_title_row,
            column=1,
            value="Full Raw Spectrum",
        ).font = Font(
            name="Calibri",
            size=12,
            bold=True,
        )

        ws.cell(
            row=data_title_row,
            column=4,
            value="Processed Data and Fit",
        ).font = Font(
            name="Calibri",
            size=12,
            bold=True,
        )

        write_dataframe(
            ws,
            raw_export,
            header_row=data_header_row,
            start_col=1,
        )

        write_dataframe(
            ws,
            fitted,
            header_row=data_header_row,
            start_col=4,
        )

        raw_last_row = data_header_row + len(raw_export)
        fit_last_row = data_header_row + len(fitted)

        fit_columns = {
            name: 4 + index
            for index, name in enumerate(fitted.columns)
        }

        # Chart 1: Full raw spectrum and estimated baseline.
        raw_chart = create_chart(
            f"{sample_id}: Raw Spectrum and Baseline"
        )

        add_series(
            raw_chart,
            ws,
            x_col=1,
            y_col=2,
            first_row=first_data_row,
            last_row=raw_last_row,
            title="Raw Spectrum",
            color="333333",
        )

        add_series(
            raw_chart,
            ws,
            x_col=fit_columns["raman_shift"],
            y_col=fit_columns["baseline"],
            first_row=first_data_row,
            last_row=fit_last_row,
            title="Baseline",
            color="E67E22",
        )

        ws.add_chart(raw_chart, "A1")

        # Chart 2: Baseline-corrected spectrum and fitted peaks.
        fit_chart = create_chart(
            f"{sample_id}: Peak Fitting"
        )
        fit_chart.x_axis.scaling.min = fit_min
        fit_chart.x_axis.scaling.max = fit_max

        curve_specs = [
            ("corrected_intensity", "Corrected Spectrum", "333333"),
            ("total_fit", "Total Fit", "E74C3C"),
            ("d_peak", "D Peak", "2980B9"),
            ("g_peak", "G Peak", "27AE60"),
            ("2D_peak", "2D Peak", "8E44AD"),
        ]

        for column_name, label, color in curve_specs:
            if column_name in fit_columns:
                add_series(
                    fit_chart,
                    ws,
                    x_col=fit_columns["raman_shift"],
                    y_col=fit_columns[column_name],
                    first_row=first_data_row,
                    last_row=fit_last_row,
                    title=label,
                    color=color,
                )

        ws.add_chart(fit_chart, "A28")

        # Chart 3: Residual.
        residual_chart = create_chart(
            f"{sample_id}: Residual",
            y_title="Residual (a.u.)",
        )
        residual_chart.x_axis.scaling.min = fit_min
        residual_chart.x_axis.scaling.max = fit_max
        residual_chart.legend = None

        add_series(
            residual_chart,
            ws,
            x_col=fit_columns["raman_shift"],
            y_col=fit_columns["residual"],
            first_row=first_data_row,
            last_row=fit_last_row,
            title="Residual",
            color="7F8C8D",
        )

        ws.add_chart(residual_chart, "A55")

        # Include the sample's summary parameters beside the data.
        sample_record = export_summary.loc[
            export_summary["Sample_ID"] == sample_id
        ].iloc[0]

        parameter_table = pd.DataFrame({
            "Parameter": sample_record.index,
            "Value": sample_record.values,
        })

        write_dataframe(
            ws,
            parameter_table,
            header_row=data_header_row,
            start_col=15,
        )

        for column_index in range(1, 17):
            ws.column_dimensions[
                get_column_letter(column_index)
            ].width = 22

        ws.column_dimensions["C"].width = 4
        ws.column_dimensions["O"].width = 26
        ws.column_dimensions["P"].width = 30


    # ---------- Apply chart formatting ----------

    from openpyxl.chart.layout import Layout, ManualLayout
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import (
        CharacterProperties,
        Font as DrawingFont,
        Paragraph,
        ParagraphProperties,
        RegularTextRun,
    )

    def text_properties(size=14, superscript=False):
        return CharacterProperties(
            latin=DrawingFont(typeface="Arial"),
            ea=DrawingFont(typeface="Arial"),
            cs=DrawingFont(typeface="Arial"),
            sz=int(size * 100),
            b=False,
            i=False,
            solidFill="000000",
            baseline=30000 if superscript else 0,
        )


    def axis_text(size=14):
        return RichText(
            p=[
                Paragraph(
                    pPr=ParagraphProperties(
                        defRPr=text_properties(size)
                    ),
                    r=[
                        RegularTextRun(
                            rPr=text_properties(size),
                            t="",
                        )
                    ],
                    endParaRPr=text_properties(size),
                )
            ]
        )


    def format_axis_title(axis, parts):
        axis.title = "".join(text for text, _ in parts)

        # Set the default title font and disable bold.
        axis.title.txPr = axis_text(14)

        # Explicitly format every text run.
        axis.title.tx.rich.p = [
            Paragraph(
                pPr=ParagraphProperties(
                    defRPr=text_properties(14)
                ),
                r=[
                    RegularTextRun(
                        rPr=text_properties(
                            size=10 if superscript else 14,
                            superscript=superscript,
                        ),
                        t=text,
                    )
                    for text, superscript in parts
                ],
                endParaRPr=text_properties(14),
            )
        ]


    def format_raman_chart(chart):
        # Chart dimensions in centimeters.
        chart.width = 12.7
        chart.height = 12.7

        # Hide the chart title and show a regular Arial 14 pt legend.
        chart.title = None
        from openpyxl.chart.legend import Legend
        chart.legend = Legend()
        chart.legend.txPr = axis_text(14)

        # Request a 4 x 4 inch plot area inside the 5 x 5 inch chart.
        chart.layout = Layout(
            manualLayout=ManualLayout(
                layoutTarget="inner",
                xMode="edge",
                yMode="edge",
                wMode="factor",
                hMode="factor",
                x=0.18,
                y=0.02,
                w=0.80,
                h=0.80,
            )
        )

        # White chart background without an outer border.
        chart.graphical_properties = GraphicalProperties(
            solidFill="FFFFFF"
        )
        chart.graphical_properties.ln.noFill = True

        # White plot area with a solid black 1 pt border.
        chart.plot_area.spPr = GraphicalProperties(
            solidFill="FFFFFF"
        )
        chart.plot_area.spPr.ln.solidFill = "000000"
        chart.plot_area.spPr.ln.w = 12700
        chart.plot_area.spPr.ln.prstDash = "solid"

        # Remove gridlines and format both axes.
        for axis in (chart.x_axis, chart.y_axis):
            axis.majorGridlines = None
            axis.minorGridlines = None
            axis.txPr = axis_text(14)
            axis.delete = False

            axis.spPr = GraphicalProperties()
            axis.spPr.ln.solidFill = "000000"
            axis.spPr.ln.w = 12700
            axis.spPr.ln.prstDash = "solid"

        # Horizontal axis: 1000 to 3000, with 500-unit intervals.
        chart.x_axis.axPos = "b"
        chart.x_axis.scaling.min = fit_min
        chart.x_axis.scaling.max = fit_max
        chart.x_axis.majorUnit = 500
        chart.x_axis.majorTickMark = "in"
        chart.x_axis.minorTickMark = "none"
        chart.x_axis.tickLblPos = "low"
        chart.x_axis.numFmt = "0"
        chart.x_axis.crossesAt = None
        chart.x_axis.crosses = "min"

        # Hide vertical-axis numbers and tick marks.
        chart.y_axis.axPos = "l"
        chart.y_axis.delete = False
        chart.y_axis.tickLblPos = "none"
        chart.y_axis.numFmt = ";;;"
        chart.y_axis.majorTickMark = "none"
        chart.y_axis.minorTickMark = "none"
        chart.y_axis.majorGridlines = None
        chart.y_axis.minorGridlines = None
        chart.y_axis.crossesAt = None
        chart.y_axis.crosses = "min"

        # Horizontal title with a superscript exponent.
        format_axis_title(
            chart.x_axis,
            [
                ("Raman shift (cm", False),
                ("-1", True),
                (")", False),
            ],
        )

        format_axis_title(
            chart.y_axis,
            [("Intensity (a.u.)", False)],
        )


    # Apply formatting to every chart.
    for worksheet in wb.worksheets:
        for chart in worksheet._charts:
            format_raman_chart(chart)

    # Set vertical limits for the normalized overview chart.
    overview_chart.y_axis.scaling.min = -0.5
    overview_chart.y_axis.scaling.max = len(overview_chart.series) + 0.5

    # Match legend order to the stacked spectra, from top to bottom.
    overview_chart.series.reverse()
    for index, series in enumerate(overview_chart.series):
        series.idx = index
        series.order = index

    # A 1.5-inch wide, 3.5-inch high legend in a 5-inch square chart.
    # Shift 0.5 inch left of the plot edge; center vertically in the plot.
    overview_chart.legend.overlay = True
    overview_chart.legend.layout = Layout(
        manualLayout=ManualLayout(
            xMode="edge", yMode="edge",
            wMode="factor", hMode="factor",
            x=0.08, y=0.07, w=0.30, h=0.70,
        )
    )

    # ---------- Save workbook ----------

    wb.save(excel_path)

    from chart_textbox import add_ratio_textbox
    ratio_lookup = {file_ids[r["file"]]: r["I_D/I_G"]
                    for r in records if r["status"] == "success"}
    add_ratio_textbox(excel_path, [ratio_lookup[s.tx.v] for s in overview_chart.series])

    print("\nExport completed.")
    print(f"Workbook saved to: {excel_path}")

    print("\nSample mapping:")
    for source_name, sample_id in file_ids.items():
        print(f"{sample_id}: {source_name}")

    print(export_summary.to_string(index=False))

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--excel-dir", type=Path, default=DEFAULT_EXCEL_DIR)
    args = parser.parse_args()

    input_dir = args.input_dir.expanduser().resolve()
    if not input_dir.is_dir():
        parser.error(f"Input directory does not exist: {input_dir}")
    files = sorted(
        (p for p in input_dir.iterdir()
         if p.is_file() and p.suffix.lower() == ".txt"),
        key=lambda p: p.name,
    )
    if not files:
        parser.error(f"No TXT files found in: {input_dir}")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    excel_dir = args.excel_dir.expanduser().resolve()
    excel_dir.mkdir(parents=True, exist_ok=True)
    excel_path = excel_dir / f"Raman_Analysis_{timestamp}.xlsx"

    model, initial_params = build_model()
    all_curve_data = {}
    records = []
    print(f"Input directory: {input_dir}")
    print(f"Files found: {len(files)}")
    for file_path in files:
        try:
            record = analyze_one_file(
                file_path, model, initial_params, all_curve_data,
                FIT_MIN, FIT_MAX, BASELINE_LAM, BASELINE_P,
            )
            records.append(record)
            print(f"Success: {file_path.name}")
        except Exception as error:
            records.append({
                "file": file_path.name,
                "status": "failed",
                "error": str(error),
            })
            print(f"Failed: {file_path.name}: {error}")

    export_workbook(
        records, all_curve_data, files, excel_path, FIT_MIN, FIT_MAX
    )
    return 1 if any(r["status"] == "failed" for r in records) else 0


if __name__ == "__main__":
    raise SystemExit(main())
